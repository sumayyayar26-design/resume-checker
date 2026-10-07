"""ATS Resume Checker - Streamlit + Google Gemini Flash.

Upload a resume (PDF, DOCX or TXT), optionally paste a job description,
and get an ATS score with concrete improvements.
"""

import io
import json
import os
import re

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

# ----------------------------------------------------------------------------
# Config
# ----------------------------------------------------------------------------
DEFAULT_MODEL = "gemini-3.5-flash"
FALLBACK_MODEL = "gemini-2.5-flash"
MAX_RESUME_CHARS = 30_000
MAX_JD_CHARS = 10_000
MIN_RESUME_CHARS = 150  # below this the file is probably a scanned image

PROMPT = """You are an expert ATS (Applicant Tracking System) analyst and senior recruiter.
Evaluate the resume below{jd_clause}.

Scoring rubric (each 0-100):
- keywords: relevant skills/keywords{jd_keywords_note}
- formatting: ATS-friendly structure (standard headings, no tables/columns/graphics, clean bullets)
- content: impact, quantified achievements, action verbs, clarity
- completeness: contact info, summary, experience, education, skills sections
Overall score = weighted average (keywords 30%, formatting 20%, content 30%, completeness 20%).

Be honest and specific. Do not invent facts that are not in the resume.
Return ONLY valid JSON with exactly this structure:
{{
  "overall_score": <int 0-100>,
  "breakdown": {{
    "keywords": <int>, "formatting": <int>, "content": <int>, "completeness": <int>
  }},
  "summary": "<2-3 sentence overall assessment>",
  "strengths": ["<string>", ...],
  "missing_keywords": ["<string>", ...],
  "improvements": [
    {{"priority": "High|Medium|Low", "section": "<resume section>",
      "issue": "<what is wrong>", "suggestion": "<concrete fix>"}}
  ],
  "rewrite_examples": [
    {{"before": "<original bullet from resume>", "after": "<improved bullet>"}}
  ]
}}

{jd_block}RESUME:
\"\"\"
{resume}
\"\"\"
"""


# ----------------------------------------------------------------------------
# Text extraction
# ----------------------------------------------------------------------------
def extract_text(filename: str, data: bytes) -> str:
    """Return plain text from a PDF, DOCX or TXT upload."""
    name = filename.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("This PDF is password-protected.")
        pages = [(page.extract_text() or "") for page in reader.pages]
        return "\n".join(pages).strip()
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:  # resumes often hide content in tables
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts).strip()
    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore").strip()
    raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")


# ----------------------------------------------------------------------------
# Gemini
# ----------------------------------------------------------------------------
def get_api_key(sidebar_key: str = "") -> str:
    """Priority: sidebar input > Streamlit secrets > environment variable."""
    if sidebar_key:
        return sidebar_key.strip()
    try:
        if "GEMINI_API_KEY" in st.secrets:
            return str(st.secrets["GEMINI_API_KEY"]).strip()
    except Exception:
        pass  # no secrets file locally - that's fine
    return os.environ.get("GEMINI_API_KEY", "").strip()


def parse_json_response(text: str) -> dict:
    """Parse model output into a dict, tolerating markdown fences/extra text."""
    text = (text or "").strip()
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.IGNORECASE)
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", text, flags=re.DOTALL)
        if match:
            return json.loads(match.group(0))
        raise ValueError("The AI response was not valid JSON. Please try again.")


def _clamp(value, default=0) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return default


def normalize_result(raw: dict) -> dict:
    """Make sure every field exists and has the right type before display."""
    breakdown = raw.get("breakdown") or {}
    keys = ["keywords", "formatting", "content", "completeness"]
    clean_breakdown = {k: _clamp(breakdown.get(k)) for k in keys}
    overall = raw.get("overall_score")
    if overall is None:
        weights = {"keywords": 0.3, "formatting": 0.2, "content": 0.3, "completeness": 0.2}
        overall = sum(clean_breakdown[k] * w for k, w in weights.items())

    def str_list(v):
        return [str(x) for x in v] if isinstance(v, list) else []

    improvements = []
    for item in raw.get("improvements") or []:
        if isinstance(item, dict):
            improvements.append(
                {
                    "priority": str(item.get("priority", "Medium")).title(),
                    "section": str(item.get("section", "General")),
                    "issue": str(item.get("issue", "")),
                    "suggestion": str(item.get("suggestion", "")),
                }
            )
    rewrites = [
        {"before": str(r.get("before", "")), "after": str(r.get("after", ""))}
        for r in (raw.get("rewrite_examples") or [])
        if isinstance(r, dict)
    ]
    return {
        "overall_score": _clamp(overall),
        "breakdown": clean_breakdown,
        "summary": str(raw.get("summary", "")),
        "strengths": str_list(raw.get("strengths")),
        "missing_keywords": str_list(raw.get("missing_keywords")),
        "improvements": improvements,
        "rewrite_examples": rewrites,
    }


def build_prompt(resume: str, job_description: str = "") -> str:
    jd = job_description.strip()[:MAX_JD_CHARS]
    return PROMPT.format(
        jd_clause=" against the job description provided" if jd else " for general ATS-friendliness",
        jd_keywords_note=" matched against the job description" if jd else " for the candidate's field",
        jd_block=f'JOB DESCRIPTION:\n"""\n{jd}\n"""\n\n' if jd else "",
        resume=resume.strip()[:MAX_RESUME_CHARS],
    )


def analyze_resume(resume: str, job_description: str, api_key: str, model: str) -> dict:
    client = genai.Client(api_key=api_key)
    config = types.GenerateContentConfig(
        temperature=0.2,
        response_mime_type="application/json",
    )
    prompt = build_prompt(resume, job_description)

    last_error = None
    for model_name in dict.fromkeys([model, FALLBACK_MODEL]):  # de-duplicated, ordered
        try:
            response = client.models.generate_content(
                model=model_name, contents=prompt, config=config
            )
            return normalize_result(parse_json_response(response.text))
        except Exception as exc:  # noqa: BLE001 - surface any API error to the UI
            last_error = exc
    raise RuntimeError(f"Gemini request failed: {last_error}")


# ----------------------------------------------------------------------------
# UI
# ----------------------------------------------------------------------------
def score_color(score: int) -> str:
    return "#16a34a" if score >= 75 else "#f59e0b" if score >= 50 else "#dc2626"


def render_results(result: dict) -> None:
    score = result["overall_score"]
    st.markdown(
        f"""
        <div style="text-align:center;padding:1.2rem;border-radius:16px;
                    border:2px solid {score_color(score)};margin-bottom:1rem;">
            <div style="font-size:0.95rem;opacity:0.7;">Your ATS Score</div>
            <div style="font-size:3.5rem;font-weight:800;color:{score_color(score)};">{score}/100</div>
        </div>
        """,
        unsafe_allow_html=True,
    )
    if result["summary"]:
        st.info(result["summary"])

    st.subheader("Score breakdown")
    cols = st.columns(4)
    labels = {
        "keywords": "Keywords",
        "formatting": "Formatting",
        "content": "Content",
        "completeness": "Completeness",
    }
    for col, (key, label) in zip(cols, labels.items()):
        value = result["breakdown"][key]
        col.metric(label, f"{value}/100")
        col.progress(value / 100)

    left, right = st.columns(2)
    with left:
        st.subheader("Strengths")
        if result["strengths"]:
            for s in result["strengths"]:
                st.markdown(f"- {s}")
        else:
            st.write("No strengths reported.")
    with right:
        st.subheader("Missing keywords")
        if result["missing_keywords"]:
            st.markdown(" ".join(f"`{k}`" for k in result["missing_keywords"]))
        else:
            st.write("No major keyword gaps found.")

    st.subheader("Recommended improvements")
    order = {"High": 0, "Medium": 1, "Low": 2}
    icons = {"High": "🔴", "Medium": "🟡", "Low": "🟢"}
    for item in sorted(result["improvements"], key=lambda i: order.get(i["priority"], 3)):
        icon = icons.get(item["priority"], "⚪")
        with st.expander(f"{icon} {item['priority']} - {item['section']}"):
            st.markdown(f"**Issue:** {item['issue']}")
            st.markdown(f"**Fix:** {item['suggestion']}")

    if result["rewrite_examples"]:
        st.subheader("Example rewrites")
        for ex in result["rewrite_examples"]:
            st.markdown(f"**Before:** {ex['before']}")
            st.markdown(f"**After:** {ex['after']}")
            st.divider()

    st.download_button(
        "Download report (JSON)",
        data=json.dumps(result, indent=2),
        file_name="ats_report.json",
        mime="application/json",
    )


def main() -> None:
    st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="wide")
    st.title("📄 ATS Resume Checker")
    st.caption("Upload your resume to get an ATS score and tips to improve it. Powered by Google Gemini.")

    with st.sidebar:
        st.header("Settings")
        sidebar_key = st.text_input(
            "Gemini API key",
            type="password",
            help="Leave empty if the key is set in Streamlit secrets or the GEMINI_API_KEY env var.",
        )
        model = st.text_input("Model", value=DEFAULT_MODEL)
        st.markdown("[Get a free API key](https://aistudio.google.com/apikey)")
        st.caption("Your resume is sent to the Gemini API for analysis and is not stored by this app.")

    uploaded = st.file_uploader("Upload your resume", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Job description (optional, but gives a much more accurate score)",
        height=160,
        placeholder="Paste the job posting here to match keywords against it...",
    )

    if st.button("Analyze resume", type="primary", disabled=uploaded is None):
        api_key = get_api_key(sidebar_key)
        if not api_key:
            st.error("Please add your Gemini API key in the sidebar (or in Streamlit secrets).")
            st.stop()
        try:
            text = extract_text(uploaded.name, uploaded.getvalue())
        except Exception as exc:  # noqa: BLE001
            st.error(f"Could not read the file: {exc}")
            st.stop()
        if len(text) < MIN_RESUME_CHARS:
            st.error(
                "Very little text was found. If your resume is a scanned image, "
                "export it as a text-based PDF or DOCX and try again."
            )
            st.stop()
        with st.spinner("Analyzing your resume..."):
            try:
                st.session_state["result"] = analyze_resume(text, job_description, api_key, model.strip() or DEFAULT_MODEL)
            except Exception as exc:  # noqa: BLE001
                st.error(str(exc))
                st.stop()

    if "result" in st.session_state:
        render_results(st.session_state["result"])


if __name__ == "__main__":
    main()
